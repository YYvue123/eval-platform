"""Build the operator manual from its maintained Markdown source."""
from pathlib import Path
import re
import json
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / '大模型智能评测平台用户使用手册.md'
OUTPUT = ROOT / '大模型智能评测平台用户使用手册_V1.0.docx'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin = section.bottom_margin = Mm(20)
section.left_margin = section.right_margin = Mm(21)
section.header_distance = section.footer_distance = Mm(9)

def font(style, size, bold=False, color='20252B'):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

font(doc.styles['Normal'], 10.5)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.22
normal.space_after = Pt(6)
normal.widow_control = True
font(doc.styles['Title'], 25, True, '000000')
font(doc.styles['Subtitle'], 12, False, '525D68')
for name, size in [('Heading 1', 17), ('Heading 2', 12.5), ('Heading 3', 11)]:
    font(doc.styles[name], size, True, '203B50')
    pf = doc.styles[name].paragraph_format
    pf.space_before = Pt(15)
    pf.space_after = Pt(8)
    pf.keep_with_next = True
    pf.keep_together = True
font(doc.styles['List Bullet'], 10.5)
font(doc.styles['List Number'], 10.5)
code_style = doc.styles.add_style('Manual Code', 1)
font(code_style, 8.5, False, '243541')
code_style.font.name = 'Consolas'
code_style.paragraph_format.space_after = Pt(0)
code_style.paragraph_format.line_spacing = 1.1
code_style.paragraph_format.left_indent = Mm(3)
code_style.paragraph_format.right_indent = Mm(3)

def field(paragraph, name):
    r = paragraph.add_run()
    el = OxmlElement('w:fldSimple')
    el.set(qn('w:instr'), name)
    r._r.addnext(el)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('用户使用手册  ·  ')
field(footer, 'PAGE')
for r in footer.runs:
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string('6C7680')

def bookmark(p, index):
    start, end = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
    start.set(qn('w:id'), str(index))
    start.set(qn('w:name'), f'chapter_{index}')
    end.set(qn('w:id'), str(index))
    p._p.insert(0, start)
    p._p.append(end)

def toc_link(label, index):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    h = OxmlElement('w:hyperlink')
    h.set(qn('w:anchor'), f'chapter_{index}')
    r, props, color, sz, text = [OxmlElement(x) for x in ['w:r','w:rPr','w:color','w:sz','w:t']]
    color.set(qn('w:val'), '203B50')
    sz.set(qn('w:val'), '21')
    props.extend([color, sz]); r.append(props)
    text.text = label; r.append(text); h.append(r); p._p.append(h)

def make_table(rows):
    cols = len(rows[0])
    t = doc.add_table(rows=0, cols=cols)
    t.autofit = False
    t.style = 'Table Grid'
    weights = {2: [0.35,0.65],3:[0.23,0.32,0.45],4:[0.19,0.27,0.27,0.27]}.get(cols,[1/cols]*cols)
    for c,w in zip(t.columns,weights): c.width = Mm(168*w)
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        trpr = t.rows[-1]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
        for j, value in enumerate(row):
            cells[j].width = Mm(168*weights[j])
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.13
            p.paragraph_format.keep_with_next = i == 0
            r = p.add_run(value.strip())
            r.font.size = Pt(9)
            r.bold = i == 0
            if i == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'),'E9EFF3')
                cells[j]._tc.get_or_add_tcPr().append(shade)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

lines = SOURCE.read_text(encoding='utf-8').splitlines()
doc.add_paragraph('大模型智能评测平台', 'Title')
doc.add_paragraph('用户使用手册', 'Title')
doc.add_paragraph('面向评测操作人员 审核人员 资源管理员与运维人员', 'Subtitle')
doc.add_paragraph('版本 1.0\n2026年9月29日')
doc.add_paragraph('从数据准备到真实评测与报告交付', 'Heading 2')
doc.add_paragraph('详细介绍各业务模块的操作入口、填写方法、状态判读和异常处理，重点说明 Agent、Skill、MCP、工具与记忆的协作方式和当前使用边界。')
doc.add_paragraph('首次使用：第1至3章\n日常评测：第4至9章\n工具和组合能力：第10至13章\nAgent与记忆：第14至19章\n专项业务与系统管理：第20至27章\n权限与实现边界：附录A至C')
doc.add_page_break()
doc.add_paragraph('目录', 'Heading 1')
headings = [x[3:] for x in lines if x.startswith('## ')]
for i,h in enumerate(headings,1): toc_link(h,i)
doc.add_page_break()

index = 0
i = next(n for n,s in enumerate(lines) if s.startswith('## '))
while i < len(lines):
    line = lines[i]
    if not line.strip(): i += 1; continue
    if line.startswith('## '):
        index += 1
        p = doc.add_paragraph(line[3:], 'Heading 1')
        # Every main chapter starts cleanly; subsections remain flowing.
        if index > 1: p.paragraph_format.page_break_before = True
        bookmark(p,index)
    elif line.startswith('### '):
        doc.add_paragraph(line[4:], 'Heading 2')
    elif line.startswith('```'):
        i += 1
        block=[]
        while i<len(lines) and not lines[i].startswith('```'):
            block.append(lines[i]); i+=1
        for k,code in enumerate(block):
            p=doc.add_paragraph(code, 'Manual Code')
            p.paragraph_format.keep_with_next = k < len(block)-1
            shading=OxmlElement('w:shd'); shading.set(qn('w:fill'),'F3F5F7')
            p._p.get_or_add_pPr().append(shading)
        doc.add_paragraph().paragraph_format.space_after=Pt(2)
    elif line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            cells=lines[i].strip('|').split('|')
            if not all(re.fullmatch(r'[ :\-]+',c) for c in cells): rows.append(cells)
            i+=1
        make_table(rows); continue
    elif line.startswith('- '):
        doc.add_paragraph(line[2:],'List Bullet')
    elif re.match(r'^\d+\. ',line):
        # Preserve source step numbers instead of a shared Word list counter.
        p=doc.add_paragraph(line)
        p.paragraph_format.left_indent=Mm(5)
        p.paragraph_format.first_line_indent=Mm(-5)
    else: doc.add_paragraph(line)
    i+=1

doc.core_properties.title='大模型智能评测平台用户使用手册'
doc.core_properties.subject='业务操作与Agent Skill MCP 工具和记忆协作'
doc.core_properties.author='评测平台项目'
doc.core_properties.keywords='评测平台,用户手册,Agent,Skill,MCP,记忆'
doc.save(OUTPUT)
# Validate archive integrity, content coverage, navigation and table shape.
from zipfile import ZipFile
from lxml import etree
with ZipFile(OUTPUT) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith('.xml'): etree.fromstring(z.read(name))
check=Document(OUTPUT)
assert len([p for p in check.paragraphs if p.style.name=='Heading 1'])==len(headings)+1
assert all(len(row.cells)==len(t.columns) for t in check.tables for row in t.rows)
report={'source_chars':len(SOURCE.read_text(encoding='utf-8')),'main_sections':len(headings),'subsections':sum(s.startswith('### ') for s in lines),'tables':len(check.tables),'paragraphs':len(check.paragraphs),'output':str(OUTPUT),'visual_render':'pending'}
(ROOT/'source-notes'/'structure-check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
