"""评测报告归档：分格式状态、证据 ID、可恢复生成。"""
from __future__ import annotations

import csv
import io
import uuid
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.models import EvalTask, ReportJob
from app.utils.jsonutil import dumps, loads


def _dest() -> Path:
    p = Path(settings.UPLOAD_DIR) / "reports"
    p.mkdir(parents=True, exist_ok=True)
    return p


def _clip(value, limit: int = 400) -> str:
    text = str(value or "").strip()
    return text if len(text) <= limit else text[:limit] + "…"


def _iso(value) -> str:
    if not value:
        return ""
    return value.isoformat() + "Z" if hasattr(value, "isoformat") else str(value)


def result_row_for_report(row) -> dict:
    """把样本结果收成报告行。已标记未评分的样本不把 0 分写成成绩。"""
    def get(key, default=None):
        if isinstance(row, dict):
            return row.get(key, default)
        return getattr(row, key, default)

    score_status = get("score_status") or ""
    scored = score_status == "scored" if score_status else get("score") is not None
    return {
        "id": get("id"),
        "item_no": get("item_no"),
        "input": _clip(get("input_content"), 500),
        "output": _clip(get("model_output"), 800),
        "reference": _clip(get("reference_answer"), 500),
        "score": get("score") if scored else None,
        "passed": bool(get("passed")) if scored else False,
        "score_status": score_status or ("scored" if scored else ""),
        "execution_status": get("execution_status") or "",
        "error_message": _clip(get("error_message"), 300),
        "latency_ms": get("latency_ms") or 0,
        "finish_reason": get("finish_reason") or "",
        "simulation": bool(get("simulation")),
    }


def _evidence_conclusions(task: EvalTask, extra: dict | None = None) -> list[dict]:
    """每条结论绑定 evidence_id，禁止无证据断言。"""
    extra = extra or {}
    conclusions = []
    conclusions.append({
        "evidence_id": f"ev-{task.id}-summary",
        "claim": "任务汇总通过率与平均分",
        "pass_rate": task.pass_rate,
        "avg_score": task.avg_score,
        "result_ref": f"task:{task.id}",
    })
    for r in (extra.get("results") or [])[:200]:
        rid = r.get("id") or r.get("item_no") or uuid.uuid4().hex[:8]
        conclusions.append({
            "evidence_id": f"ev-{task.id}-r{rid}",
            "claim": "样本评分",
            "score": r.get("score"),
            "passed": r.get("passed"),
            "result_ref": f"result:{rid}",
        })
    return conclusions


def _payload(task: EvalTask, extra: dict | None = None) -> dict:
    extra = extra or {}
    evidence = _evidence_conclusions(task, extra)
    results = []
    for row in extra.get("results") or []:
        results.append(row if isinstance(row, dict) and "input" in row else result_row_for_report(row))
    body = {
        "task_id": task.id,
        "name": task.name,
        "status": task.status,
        "task_type": getattr(task, "task_type", "") or "",
        "scene": task.scene,
        "industry": task.industry,
        "template_code": getattr(task, "template_code", "") or "",
        "trial_run": bool(getattr(task, "trial_run", False)),
        "simulation": bool(getattr(task, "simulation", False)),
        "priority": getattr(task, "priority", None),
        "depends_on_id": getattr(task, "depends_on_id", None),
        "dataset_id": task.dataset_id,
        "dataset_name": extra.get("dataset_name") or "",
        "dataset_version_id": task.dataset_version_id,
        "model_id": task.model_id,
        "model_name": extra.get("model_name") or "",
        "model_version_id": getattr(task, "model_version_id", None),
        "channel_type": extra.get("channel_type") or "",
        "prompt_id": task.prompt_id,
        "prompt_name": extra.get("prompt_name") or "",
        "prompt_version_id": task.prompt_version_id,
        "judge_resource_id": task.judge_resource_id,
        "tool_version": getattr(task, "tool_version", "") or "",
        "snapshot_id": task.snapshot_id,
        "batch_id": task.batch_id,
        "total": task.total,
        "success_count": task.success_count,
        "fail_count": task.fail_count,
        "skip_count": getattr(task, "skip_count", 0) or 0,
        "avg_score": task.avg_score,
        "pass_rate": task.pass_rate,
        "token_quota": getattr(task, "token_quota", 0) or 0,
        "tokens_used": getattr(task, "tokens_used", 0) or 0,
        "error_message": getattr(task, "error_message", "") or "",
        "started_at": _iso(getattr(task, "started_at", None)),
        "finished_at": _iso(getattr(task, "finished_at", None)),
        "report_summary": task.report_summary,
        "conclusions": evidence,
        "evidence_ids": [c["evidence_id"] for c in evidence],
        "results": results,
    }
    passthrough = {k: v for k, v in extra.items() if k not in body and k != "results"}
    body.update(passthrough)
    return body


def write_task_report(task: EvalTask, extra: dict | None = None) -> str:
    """同步写报告（兼容旧路径）；失败可见于 formats 元数据文件。"""
    dest = _dest()
    payload = _payload(task, extra)
    formats: dict[str, dict] = {}
    path = dest / f"task-{task.id}.json"
    try:
        path.write_text(dumps(payload), encoding="utf-8")
        formats["json"] = {"status": "ready", "path": str(path), "error": ""}
    except Exception as exc:  # noqa: BLE001
        formats["json"] = {"status": "failed", "path": "", "error": str(exc)}

    md_path = dest / f"task-{task.id}.md"
    try:
        md_path.write_text(_markdown(payload), encoding="utf-8")
        formats["md"] = {"status": "ready", "path": str(md_path), "error": ""}
    except Exception as exc:  # noqa: BLE001
        formats["md"] = {"status": "failed", "path": "", "error": str(exc)}

    writers = (
        (write_html, "html"),
        (write_csv, "csv"),
        (write_xlsx, "xlsx"),
        (write_docx, "docx"),
        (write_pdf_summary, "pdf"),
    )
    for writer, ext in writers:
        fpath = dest / f"task-{task.id}.{ext}"
        try:
            writer(fpath, payload)
            formats[ext] = {"status": "ready", "path": str(fpath), "error": ""}
        except Exception as exc:  # noqa: BLE001
            formats[ext] = {"status": "failed", "path": "", "error": str(exc)}

    meta = dest / f"task-{task.id}.formats.json"
    meta.write_text(dumps({"task_id": task.id, "formats": formats, "evidence_ids": payload.get("evidence_ids")}), encoding="utf-8")
    return str(path)


async def render_report_job(db: AsyncSession, task: EvalTask, extra: dict | None = None) -> ReportJob:
    job = ReportJob(task_id=task.id, status="rendering", formats_json="{}", evidence_json="[]")
    db.add(job)
    await db.flush()
    try:
        path = write_task_report(task, extra)
        meta_path = Path(settings.UPLOAD_DIR) / "reports" / f"task-{task.id}.formats.json"
        meta = loads(meta_path.read_text(encoding="utf-8"), {}) if meta_path.is_file() else {}
        formats = meta.get("formats") or {}
        failed = [k for k, v in formats.items() if v.get("status") == "failed"]
        job.formats_json = dumps(formats)
        job.evidence_json = dumps((extra or {}).get("conclusions") or loads(Path(path).read_text(encoding="utf-8"), {}).get("conclusions") or [])
        if failed and "json" in failed:
            job.status = "failed"
            job.error_message = "；".join(f"{k}:{formats[k].get('error')}" for k in failed)
        elif failed:
            job.status = "ready"  # 主格式成功，部分格式失败可见
            job.error_message = "partial:" + ",".join(failed)
        else:
            job.status = "ready"
            job.error_message = ""
        task.report_path = path
    except Exception as exc:  # noqa: BLE001
        job.status = "failed"
        job.error_message = str(exc)
        job.formats_json = dumps({})
    await db.flush()
    return job


def report_file(task_id: int, fmt: str) -> Path:
    ext = {"json": "json", "md": "md", "csv": "csv", "xlsx": "xlsx", "docx": "docx", "html": "html", "pdf": "pdf"}.get(fmt, "json")
    return _dest() / f"task-{task_id}.{ext}"


def report_formats_status(task_id: int) -> dict:
    meta = _dest() / f"task-{task_id}.formats.json"
    if not meta.is_file():
        # 推断存在文件
        formats = {}
        for ext in ("json", "md", "csv", "xlsx", "docx", "html", "pdf"):
            p = report_file(task_id, ext)
            formats[ext] = {"status": "ready" if p.is_file() else "missing", "path": str(p) if p.is_file() else "", "error": "" if p.is_file() else "not_generated"}
        return {"task_id": task_id, "formats": formats}
    return loads(meta.read_text(encoding="utf-8"), {"task_id": task_id, "formats": {}})


def _markdown(payload: dict) -> str:
    lines = [
        f"# 评测报告 {payload.get('name')}",
        "",
        "## 任务",
        f"- 任务 ID：{payload.get('task_id')}",
        f"- 状态：{payload.get('status')}",
        f"- 类型：{payload.get('task_type') or '-'} · 模板：{payload.get('template_code') or '-'}",
        f"- 场景：{payload.get('scene') or '-'} · 行业：{payload.get('industry') or '-'}",
        f"- 模式：{'试跑' if payload.get('trial_run') else '正式'} · 模拟：{'是' if payload.get('simulation') else '否'}",
        f"- 优先级：{payload.get('priority') if payload.get('priority') is not None else '-'} · 依赖任务：{payload.get('depends_on_id') or '-'}",
        f"- 开始：{payload.get('started_at') or '-'} · 结束：{payload.get('finished_at') or '-'}",
        "",
        "## 配置",
        f"- 数据集：{payload.get('dataset_name') or '-'}（版本 {payload.get('dataset_version_id') or '-'}）",
        f"- 模型：{payload.get('model_name') or '-'}（版本 {payload.get('model_version_id') or '-'}，通道 {payload.get('channel_type') or '-'}）",
        f"- 提示词：{payload.get('prompt_name') or '-'}（版本 {payload.get('prompt_version_id') or '-'}）",
        f"- 裁判：{payload.get('judge_resource_id') or '-'} · 工具版本：{payload.get('tool_version') or '-'}",
        f"- 批次：{payload.get('batch_id') or '-'} · 快照：{payload.get('snapshot_id') or '-'}",
        "",
        "## 结果",
        f"- 样本 {payload.get('total')} · 通过 {payload.get('success_count')} · 失败 {payload.get('fail_count')} · 跳过 {payload.get('skip_count')}",
        f"- 通过率 {payload.get('pass_rate')} · 平均分 {payload.get('avg_score')}",
        f"- Token 配额 {payload.get('token_quota')} · 已用 {payload.get('tokens_used')}",
        f"- 摘要：{payload.get('report_summary') or '-'}",
    ]
    if payload.get("error_message"):
        lines.append(f"- 错误：{payload.get('error_message')}")
    lines.extend(["", "## 证据"])
    for c in payload.get("conclusions") or []:
        lines.append(f"- `{c.get('evidence_id')}` {c.get('claim')}（{c.get('result_ref')}）")
    rows = payload.get("results") or []
    lines.extend(["", "## 样本明细", ""])
    if not rows:
        lines.append("无逐条结果。")
    else:
        lines.append("| # | 通过 | 分数 | 状态 | 输入 | 输出 | 参考 | 错误 |")
        lines.append("| --- | --- | --- | --- | --- | --- | --- | --- |")
        for r in rows[:200]:
            cells = [
                str(r.get("item_no") or ""),
                "是" if r.get("passed") else "否",
                "" if r.get("score") is None else str(r.get("score")),
                str(r.get("score_status") or r.get("execution_status") or ""),
                _md_cell(r.get("input")),
                _md_cell(r.get("output")),
                _md_cell(r.get("reference")),
                _md_cell(r.get("error_message")),
            ]
            lines.append("| " + " | ".join(cells) + " |")
        if len(rows) > 200:
            lines.append(f"\n其余 {len(rows) - 200} 条见 JSON / Excel。")
    lines.append("")
    return "\n".join(lines)


def _md_cell(value) -> str:
    return _clip(value, 120).replace("|", "/").replace("\n", " ")


def write_html(path: Path, payload: dict) -> None:
    rows = payload.get("results") or []
    body = "".join(
        "<tr>"
        f"<td>{escape(str(r.get('item_no','')))}</td>"
        f"<td>{escape('是' if r.get('passed') else '否')}</td>"
        f"<td>{escape('' if r.get('score') is None else str(r.get('score')))}</td>"
        f"<td>{escape(str(r.get('score_status') or r.get('execution_status') or ''))}</td>"
        f"<td>{escape(_clip(r.get('input'), 200))}</td>"
        f"<td>{escape(_clip(r.get('output'), 240))}</td>"
        f"<td>{escape(_clip(r.get('reference'), 200))}</td>"
        f"<td>{escape(_clip(r.get('error_message'), 160))}</td>"
        "</tr>"
        for r in rows[:300]
    )
    ev = "".join(
        f"<li><code>{escape(c.get('evidence_id',''))}</code> {escape(str(c.get('claim','')))}</li>"
        for c in (payload.get("conclusions") or [])[:50]
    )
    meta = "".join(
        f"<li>{escape(label)}：{escape(str(payload.get(key) or '-'))}</li>"
        for label, key in (
            ("状态", "status"),
            ("场景", "scene"),
            ("行业", "industry"),
            ("数据集", "dataset_name"),
            ("模型", "model_name"),
            ("通道", "channel_type"),
            ("提示词", "prompt_name"),
            ("裁判", "judge_resource_id"),
            ("通过率", "pass_rate"),
            ("平均分", "avg_score"),
            ("已用 Token", "tokens_used"),
        )
    )
    path.write_text(
        f"<!DOCTYPE html><html lang='zh-CN'><head><meta charset='utf-8'><title>报告 {escape(str(payload.get('name')))}</title></head>"
        f"<body><h1>{escape(str(payload.get('name')))}</h1>"
        f"<p>{escape(str(payload.get('report_summary') or ''))}</p>"
        f"<h2>关键信息</h2><ul>{meta}</ul>"
        f"<h2>证据</h2><ul>{ev}</ul>"
        f"<h2>样本明细</h2><table border='1'><tr><th>#</th><th>通过</th><th>分数</th><th>状态</th><th>输入</th><th>输出</th><th>参考</th><th>错误</th></tr>{body}</table></body></html>",
        encoding="utf-8",
    )


def write_csv(path: Path, payload: dict) -> None:
    summary_keys = (
        "task_id", "name", "status", "task_type", "scene", "industry", "template_code",
        "trial_run", "simulation", "dataset_name", "dataset_version_id", "model_name",
        "model_version_id", "channel_type", "prompt_name", "judge_resource_id", "tool_version",
        "total", "success_count", "fail_count", "skip_count", "pass_rate", "avg_score",
        "token_quota", "tokens_used", "batch_id", "snapshot_id", "started_at", "finished_at",
        "error_message", "report_summary",
    )
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["field", "value"])
        for key in summary_keys:
            w.writerow([key, payload.get(key)])
        w.writerow([])
        w.writerow(["evidence_id", "claim", "result_ref", "score", "passed"])
        for c in payload.get("conclusions") or []:
            w.writerow([c.get("evidence_id"), c.get("claim"), c.get("result_ref"), c.get("score"), c.get("passed")])
        w.writerow([])
        w.writerow(["item_no", "passed", "score", "score_status", "execution_status", "latency_ms", "input", "output", "reference", "error_message"])
        for r in payload.get("results") or []:
            w.writerow([
                r.get("item_no"), r.get("passed"), r.get("score"), r.get("score_status"),
                r.get("execution_status"), r.get("latency_ms"), r.get("input"), r.get("output"),
                r.get("reference"), r.get("error_message"),
            ])


def write_xlsx(path: Path, payload: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    ws.append(["字段", "值"])
    for k in (
        "task_id", "name", "status", "task_type", "scene", "industry", "template_code",
        "trial_run", "dataset_name", "model_name", "channel_type", "prompt_name",
        "judge_resource_id", "pass_rate", "avg_score", "total", "success_count", "fail_count",
        "skip_count", "tokens_used", "token_quota", "error_message", "report_summary",
        "started_at", "finished_at",
    ):
        ws.append([k, payload.get(k)])
    ws2 = wb.create_sheet("evidence")
    ws2.append(["evidence_id", "claim", "result_ref"])
    for c in payload.get("conclusions") or []:
        ws2.append([c.get("evidence_id"), c.get("claim"), c.get("result_ref")])
    ws3 = wb.create_sheet("results")
    ws3.append(["item_no", "passed", "score", "score_status", "execution_status", "latency_ms", "input", "output", "reference", "error_message"])
    for r in payload.get("results") or []:
        ws3.append([
            r.get("item_no"), r.get("passed"), r.get("score"), r.get("score_status"),
            r.get("execution_status"), r.get("latency_ms"), r.get("input"), r.get("output"),
            r.get("reference"), r.get("error_message"),
        ])
    wb.save(path)


def write_docx(path: Path, payload: dict) -> None:
    lines = [
        f"评测报告 {payload.get('name')}",
        f"状态 {payload.get('status')} 类型 {payload.get('task_type') or '-'} 场景 {payload.get('scene')} 行业 {payload.get('industry')}",
        f"数据集 {payload.get('dataset_name') or '-'} 模型 {payload.get('model_name') or '-'} 通道 {payload.get('channel_type') or '-'}",
        f"提示词 {payload.get('prompt_name') or '-'} 裁判 {payload.get('judge_resource_id') or '-'}",
        f"样本 {payload.get('total')} 通过 {payload.get('success_count')} 失败 {payload.get('fail_count')} 跳过 {payload.get('skip_count')}",
        f"通过率 {payload.get('pass_rate')} 平均分 {payload.get('avg_score')} Token {payload.get('tokens_used')}/{payload.get('token_quota')}",
        str(payload.get("report_summary") or ""),
    ]
    if payload.get("error_message"):
        lines.append(f"错误 {payload.get('error_message')}")
    for c in (payload.get("conclusions") or [])[:12]:
        lines.append(f"证据 {c.get('evidence_id')}: {c.get('claim')}")
    for r in (payload.get("results") or [])[:30]:
        lines.append(
            f"样本 {r.get('item_no')} 分数 {r.get('score')} 通过 {r.get('passed')} 输入 {_clip(r.get('input'), 80)} 输出 {_clip(r.get('output'), 80)}"
        )
    paragraphs = "".join("<w:p><w:r><w:t>" + escape(line) + "</w:t></w:r></w:p>" for line in lines if line)
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{paragraphs}</w:body></w:document>"
    )
    ctypes = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
        "</Relationships>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ctypes)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", document_xml)
    path.write_bytes(buf.getvalue())


def write_pdf_summary(path: Path, payload: dict) -> None:
    lines = [
        _pdf_esc(f"Eval report task {payload.get('task_id')} {payload.get('status')}"),
        _pdf_esc(f"scene={payload.get('scene')} industry={payload.get('industry')} trial={payload.get('trial_run')}"),
        _pdf_esc(f"dataset={payload.get('dataset_name') or payload.get('dataset_id')} model={payload.get('model_name') or payload.get('model_id')}"),
        _pdf_esc(f"judge={payload.get('judge_resource_id')} channel={payload.get('channel_type') or '-'}"),
        _pdf_esc(f"pass_rate={payload.get('pass_rate')} avg={payload.get('avg_score')} total={payload.get('total')} tokens={payload.get('tokens_used')}"),
        _pdf_esc(f"evidence={(payload.get('evidence_ids') or ['none'])[0]}"),
    ]
    body = "BT /F1 11 Tf 72 740 Td "
    body += " Tj 0 -16 Td ".join(f"({line})" for line in lines)
    body += " Tj ET\n"
    stream = body.encode("latin-1", "replace")
    parts = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(parts, 1):
        offsets.append(len(out))
        out.extend(f"{i} 0 obj\n".encode("ascii"))
        out.extend(obj)
        out.extend(b"\nendobj\n")
    xref = len(out)
    out.extend(f"xref\n0 {len(parts)+1}\n0000000000 65535 f \n".encode("ascii"))
    for off in offsets:
        out.extend(f"{off:010d} 00000 n \n".encode("ascii"))
    out.extend(
        f"trailer << /Size {len(parts)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii")
    )
    path.write_bytes(bytes(out))


def _pdf_esc(s: str) -> str:
    return str(s).replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")[:200]
