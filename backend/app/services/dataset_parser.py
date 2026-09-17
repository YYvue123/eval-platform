"""评测数据文件解析：Excel / CSV / JSON / JSONL / TXT / ZIP。"""
from __future__ import annotations

import csv
import io
import json
import zipfile
from pathlib import Path


INPUT_KEYS = ("input", "question", "prompt", "query", "input_content", "原文", "问题")
REF_KEYS = ("reference", "answer", "output", "label", "reference_answer", "expected_output", "答案", "参考答案")
REQ_KEYS = ("requirement", "task_requirement", "instruction", "任务要求")
LABEL_KEYS = ("label", "tag", "data_label", "标签")
DIFF_KEYS = ("difficulty", "difficulty_level", "难度")


def _pick(row: dict, keys: tuple[str, ...], default: str = "") -> str:
    lower = {str(k).strip().lower(): v for k, v in row.items()}
    for key in keys:
        val = lower.get(key.lower())
        if val is not None and str(val).strip() != "":
            return str(val)
    return default


def normalize_row(row: dict, index: int) -> dict:
    extra = {k: v for k, v in row.items() if k}
    return {
        "item_no": index,
        "input_content": _pick(row, INPUT_KEYS),
        "reference_answer": _pick(row, REF_KEYS),
        "expected_output": _pick(row, ("expected", "expected_output")),
        "task_requirement": _pick(row, REQ_KEYS),
        "difficulty_level": _pick(row, DIFF_KEYS),
        "data_label": _pick(row, LABEL_KEYS),
        "extended_content": extra,
    }


def parse_bytes(filename: str, data: bytes) -> list[dict]:
    name = (filename or "").lower()
    if name.endswith(".zip"):
        return _parse_zip(data)
    if name.endswith(".jsonl"):
        return _parse_jsonl(data.decode("utf-8-sig"))
    if name.endswith(".json"):
        return _parse_json(data.decode("utf-8-sig"))
    if name.endswith(".csv") or name.endswith(".txt"):
        text = data.decode("utf-8-sig")
        if name.endswith(".txt") and "," not in text.split("\n", 1)[0] and "\t" not in text.split("\n", 1)[0]:
            return _parse_txt(text)
        return _parse_csv(text)
    if name.endswith(".xlsx") or name.endswith(".xls"):
        return _parse_excel(data)
    raise ValueError("不支持的文件格式，请上传 Excel、CSV、JSON、JSONL、TXT 或 ZIP")


def _parse_json(text: str) -> list[dict]:
    payload = json.loads(text)
    if isinstance(payload, dict):
        for key in ("items", "data", "samples"):
            if isinstance(payload.get(key), list):
                payload = payload[key]
                break
        else:
            payload = [payload]
    if not isinstance(payload, list):
        raise ValueError("JSON 内容必须是对象数组")
    rows = []
    for i, item in enumerate(payload, start=1):
        if isinstance(item, str):
            item = {"input": item}
        if not isinstance(item, dict):
            continue
        rows.append(normalize_row(item, i))
    return rows


def _parse_jsonl(text: str) -> list[dict]:
    rows = []
    for i, line in enumerate(text.splitlines(), start=1):
        line = line.strip()
        if not line:
            continue
        obj = json.loads(line)
        if isinstance(obj, str):
            obj = {"input": obj}
        rows.append(normalize_row(obj, i))
    return rows


def _parse_csv(text: str) -> list[dict]:
    sample = text[:2048]
    dialect = csv.Sniffer().sniff(sample, delimiters=",\t;") if sample.strip() else csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    rows = []
    for i, row in enumerate(reader, start=1):
        rows.append(normalize_row({k: v for k, v in row.items() if k is not None}, i))
    return rows


def _parse_txt(text: str) -> list[dict]:
    rows = []
    n = 0
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        n += 1
        rows.append(normalize_row({"input": line}, n))
    return rows


def _parse_excel(data: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("服务器未安装 openpyxl，无法解析 Excel") from exc
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers = next(rows_iter, None)
    if not headers:
        return []
    keys = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(headers)]
    rows = []
    for i, values in enumerate(rows_iter, start=1):
        row = {keys[j]: ("" if v is None else v) for j, v in enumerate(values) if j < len(keys)}
        rows.append(normalize_row(row, i))
    return rows


def _parse_zip(data: bytes) -> list[dict]:
    rows = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for info in zf.infolist():
            if info.is_dir() or info.filename.startswith("__MACOSX"):
                continue
            suffix = Path(info.filename).suffix.lower()
            if suffix not in {".json", ".jsonl", ".csv", ".txt", ".xlsx", ".xls"}:
                continue
            part = zf.read(info.filename)
            rows.extend(parse_bytes(info.filename, part))
    for i, row in enumerate(rows, start=1):
        row["item_no"] = i
    return rows
